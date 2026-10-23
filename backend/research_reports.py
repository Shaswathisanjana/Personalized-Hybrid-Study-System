"""Readable report exports from existing validated claims; no AI calls or CACM writes."""
from io import BytesIO
from collections import defaultdict
from backend.report_synthesis import enrich_sources, synthesis_paragraphs
from app.research_agent.docx_exporter import SECTIONS, clean
from app.research_agent.research_report import report_sources

CONTENT_SECTIONS=set(SECTIONS)-{'title','references','limitations'}
FIELD_LABELS={
    'overview':'Overview excerpt', 'research_problem':'Problem-related excerpt',
    'objectives':'Objective-related excerpt', 'methodology':'Method-related excerpt',
    'models_algorithms':'Model-related excerpt (implemented, reviewed or proposed role not independently verified)',
    'dataset':'Dataset-related excerpt (dataset identity may be unspecified)',
    'evaluation_metrics':'Evaluation-related excerpt (protocol and metric definition may be unspecified)',
    'experimental_setup':'Experimental-setup excerpt',
    'findings':'Author-reported finding (not independently reproduced)',
    'limitations':'Limitation-related excerpt (specific limitation requires verification)',
}


def failure_summary(errors):
    text=' '.join(errors).lower()
    if any(token in text for token in ('503', 'unavailable', 'high demand', '429')):
        return 'The analysis provider was temporarily unavailable or rate limited. Retry preparation later.'
    if any(token in text for token in ('quotation', 'json', 'defined analysis fields')):
        return 'The provider response could not be verified against the source. No unsupported claims were accepted; retry preparation.'
    return 'Source processing failed. Retry preparation or exclude the source.'


def scope_label(source):
    return 'abstract-only' if source['scope']=='abstract_only' else 'full-text excerpts'

def readiness(documents):
    selected=[d for d in documents if d['record'].get('selected',True)]
    sources=report_sources(selected)
    total=sum(s['total'] for s in sources)
    completed=sum(s['completed'] for s in sources)
    claims=sum(len(items) for s in sources for items in s['claims'].values())
    pending=sum(d['record']['status'] in {'queued','extraction_error'} for d in selected)
    errors=sum(len(s['errors']) for s in sources)
    incomplete=[s['title'] for s in sources if s['completed']<s['total'] or not s['total'] or s['status']!='ready']
    ready=bool(sources and completed and claims and not pending and not errors and not incomplete)
    if not sources:message='Select papers or PDFs for the report.'
    elif not completed:message='No selected source has been analyzed. Prepare selected sources before generating an analysis report.'
    elif not claims:message='Analysis returned no supported claims. Inspect source coverage or add usable sources; only a bibliography is available.'
    elif not ready:message='Some selected sources are incomplete or failed. Prepare another batch, deselect them, or explicitly allow a labelled partial report.'
    else:message='Selected source excerpts have finished processing. Review field coverage and source context before export.'
    return dict(selected_sources=len(sources),total_chunks=total,analyzed_chunks=completed,supported_claims=claims,
        pending_extraction=pending,failed_chunks=errors,incomplete_sources=incomplete,ready=ready,
        can_export_partial=bool(completed and claims),message=message)

def validate_export(documents,sections,allow_partial=False,metadata_only=False):
    if not sections or set(sections)-set(SECTIONS):raise ValueError('Choose supported report sections.')
    if metadata_only and set(sections)&CONTENT_SECTIONS:raise ValueError('Bibliography exports cannot include analysis sections.')
    status=readiness(documents)
    if not status['selected_sources']:raise ValueError('Select at least one research source.')
    if not metadata_only and not status['ready']:
        if not (allow_partial and status['can_export_partial']):raise ValueError(status['message'])
    return status

def report_blocks(title,topic,sources,sections,partial=False,metadata_only=False):
    """A shared document model keeps Word and Markdown contents consistent."""
    sources=enrich_sources(sources)
    chosen=set(sections)
    if not chosen or chosen-set(SECTIONS):raise ValueError('Choose supported report sections.')
    if not sources:raise ValueError('Select report sources.')
    blocks=[]
    def h(text,level=1):blocks.append(('heading',(text,level)))
    def p(text):blocks.append(('paragraph',str(text)))
    def table(headers,rows):blocks.append(('table',(headers,rows)))
    indexes={s['id']:i for i,s in enumerate(sources,1)}
    def cite(source,claim):
        where=('abstract-only' if claim['page'] is None else 'page '+str(claim['page']))
        return '['+str(indexes[source['id']])+', '+where+', chunk '+claim['chunk_id']+']'
    def evidence(source,fields):
        found=False
        for field in fields:
            for claim in source['claims'][field]:
                if claim.get('observation') and claim['observation']!=claim['quote']:
                    p('Saved AI interpretation; check against the quotation: '+claim['observation'])
                p(FIELD_LABELS.get(field,field.replace('_',' ').capitalize())+': "'+claim['quote']+'" '+cite(source,claim));found=True
        if not found:
            p('Not available: '+('this source has not been analyzed.' if not source['completed'] else 'not reported in the analyzed source excerpts.'))
    if metadata_only:h('Source bibliography',0)
    elif 'title' in chosen:h(title.strip() or 'Research Analysis Report',0)
    if topic.strip():p('Research topic: '+topic.strip())
    if metadata_only:
        p('References and source coverage only. This document does not present completed research analysis.')
    else:
        h('Report overview')
        counts=sum(s['completed'] for s in sources);total=sum(s['total'] for s in sources)
        abstracts=sum(s['scope']=='abstract_only' for s in sources)
        p(f'This report covers {len(sources)} selected sources: {abstracts} abstract-only sources and {len(sources)-abstracts} sources of full-text excerpts. {counts}/{total} available chunks were analyzed.')
        if partial:p('PARTIAL REPORT: selected sources have incomplete analysis or recoverable failures. Conclusions cannot be extended to unprocessed content.')
        p('Report narratives are built from attributed source quotations. Field assignments are conservatively enriched from explicit dataset sizes and metric mentions in saved evidence. Source-supported observations are presented as attributed excerpts. Gaps and future directions are labelled system suggestions. Missing data are not inferred from titles. This is not a systematic literature review.')
        p('Quotation verification confirms that the words occur in the supplied source; it does not verify scientific correctness, the AI field classification, or whether a method was implemented rather than reviewed or proposed.')
        p('Analysis completion means processing of available excerpts, not complete coverage of each paper. Reported percentages are not directly comparable without matching datasets, splits, metric definitions, thresholds and evaluation conditions. No best-performing method is established here.')
        omitted_sections=[label for key,label in SECTIONS.items() if key not in chosen]
        if omitted_sections:p('Sections not requested: '+', '.join(omitted_sections)+'. Use Select all report sections in the application for a broader report.')
        if chosen & CONTENT_SECTIONS:
            h('Evidence summary')
            p('The following counts describe saved source excerpts, not the quality or completeness of the publications.')
            for field in ('methodology','dataset','evaluation_metrics','findings','limitations'):
                covered=[s for s in sources if s['claims'][field]]
                refs=', '.join('['+str(indexes[s['id']])+']' for s in covered)
                p(field.replace('_',' ').capitalize()+f': excerpts available for {len(covered)}/{len(sources)} selected sources'+(' ('+refs+')' if refs else '')+'.')
            p('A missing excerpt means information was not extracted from the available content. It does not prove the paper lacks that information or establish a research gap.')
    h('Source coverage')
    table(['Reference','Paper','Coverage','Analyzed chunks'],[[f'[{indexes[s["id"]]}]',s['title'],scope_label(s),f'{s["completed"]}/{s["total"]}'] for s in sources])
    analyzed=[s for s in sources if s['completed']]
    omitted=[s for s in sources if not s['completed']]
    if omitted and not metadata_only:
        p('Unanalyzed sources are listed in the coverage table and references only; they are excluded from analytical comparisons and synthesis.')
    groups=dict(overviews=('overview','research_problem','findings'),objectives=('objectives',),
        methods=('methodology','models_algorithms'),evaluation=('dataset','experimental_setup','evaluation_metrics'))
    if not metadata_only:
        for key,fields in groups.items():
            if key not in chosen:continue
            h(SECTIONS[key])
            for s in analyzed:
                h(f'[{indexes[s["id"]]}] '+s['title'],2)
                p('Source coverage: '+scope_label(s)+'.')
                if key=='overviews':
                    overview=synthesis_paragraphs([s],cite,[('research_problem','the problem addressed'),('methodology','the approach described'),('findings','the reported outcome')])
                    if overview:
                        for paragraph in overview:p(paragraph)
                    else:evidence(s,fields)
                else:evidence(s,fields)
        if 'comparison' in chosen:
            h(SECTIONS['comparison'])
            rows=[]
            for s in analyzed:
                def values(field):
                    return '\n'.join(c['quote']+' '+cite(s,c) for c in s['claims'][field]) or 'Not available in analyzed excerpts' if s['completed'] else 'Not analyzed'
                rows.append([f'[{indexes[s["id"]]}] '+s['title'],scope_label(s),values('methodology'),values('dataset'),values('evaluation_metrics')])
            table(['Paper','Scope','Methodology','Dataset','Evaluation'],rows)
            p('Cells contain source quotations assigned to fields by the analyzer. Generic references to datasets do not identify a dataset; future-work model mentions do not establish implemented methods. Unavailable cells do not imply that the paper has no method or evaluation.')
            h('Reported results and comparison boundaries',2)
            for source in analyzed:
                h(f'[{indexes[source["id"]]}] '+source['title'],2)
                evidence(source,('findings',))
            p('These are separately reported results, not a common benchmark. Accuracy, mAP, FPS and parameter counts measure different things; no numerical ranking or cross-paper improvement is inferred.')
        if 'themes' in chosen:
            h(SECTIONS['themes'])
            p('This synthesis connects attributed evidence across the selected papers. Quoted statements remain author reports; the connecting prose does not claim agreement, causation or common evaluation conditions.')
            narrative=synthesis_paragraphs(analyzed,cite,[('research_problem','the research problem'),('methodology','the approaches described'),('dataset','the available dataset descriptions'),('findings','the reported outcomes'),('limitations','the reported limitations')])
            for paragraph in narrative:p(paragraph)
            if not narrative:p('No source-supported analysis is available.')
        if 'differences' in chosen:
            h(SECTIONS['differences'])
            p('Similarities below mean overlap in available evidence categories, not proof of equivalent algorithms. Differences describe separately attributed source accounts; absent evidence is not a negative finding.')
            for field,label in [('methodology','method descriptions'),('dataset','dataset descriptions'),('evaluation_metrics','evaluation descriptions'),('findings','reported outcomes')]:
                covered=[source for source in analyzed if source['claims'][field]]
                if not covered:continue
                h(label.capitalize(),2)
                p(f'{len(covered)} of {len(analyzed)} analyzed sources provide {label}.')
                for paragraph in synthesis_paragraphs(covered,cite,[(field,label)]):p(paragraph)
                if len(covered)>1:p('These source accounts can be examined side by side, but their wording alone does not establish equivalent methods or comparable outcomes.')
        for key in ('gaps','directions'):
            if key not in chosen:continue
            h(SECTIONS[key]);p('System-generated research suggestions, not paper findings or verified novelty claims. A broader literature review and feasibility checks are required.')
            anchors=[(source,claim) for source in analyzed for claim in source['claims']['limitations']]
            if not anchors:
                p('No suggestion is generated because no source-supported limitation was extracted. Missing report fields alone are not evidence of a research gap.')
            for source,claim in anchors:
                p('Source excerpt requiring contextual verification: "'+claim['quote']+'" '+cite(source,claim))
                if key=='gaps':
                    p('Suggested verification question: Does this passage describe a specific unresolved limitation, and does that limitation remain unresolved in later literature? This report does not establish either conclusion.')
                else:
                    p('Suggested next step: inspect the full paper for the context of this passage, identify a measurable research question if warranted, and check later literature before designing an experiment. No particular dataset, method or expected improvement is proposed without supporting evidence.')
    if 'limitations' in chosen:
        h(SECTIONS['limitations'])
        for s in sources:
            h(f'[{indexes[s["id"]]}] '+s['title'],2)
            p(f'{scope_label(s)}; {s["completed"]}/{s["total"]} chunks analyzed.')
            for warning in s['warnings']:p(warning)
            if s['errors']:
                p('Analysis incomplete: '+failure_summary(s['errors'])+' Detailed diagnostics remain in the application.')
            if not metadata_only:evidence(s,('limitations',))
    # References remain mandatory for traceable evidence, even when the optional
    # reference section is unchecked. Omit unavailable metadata rather than invent it.
    h('References and source details')
    for s in sources:
        h(f'[{indexes[s["id"]]}] '+s['title'],2)
        p('Authors: '+(', '.join(s['authors']) or 'Not supplied')+'. Year: '+str(s['year'] or 'Not supplied')+'.')
        if s.get('doi'):p('DOI: '+s['doi'].removeprefix('https://doi.org/'))
        if s.get('url'):p('Source: '+s['url'])
        if s.get('filename'):p('Uploaded file: '+s['filename'])
        p('Coverage: '+scope_label(s)+'.')
    return blocks

def readable_docx(title,topic,sources,sections,**options):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document=Document()
    normal=document.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    for name in ('Title','Heading 1','Heading 2'):
        document.styles[name].font.color.rgb=RGBColor.from_string('285B44')
        document.styles[name].paragraph_format.keep_with_next=True
    for section in document.sections:
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.7)
    document.sections[0].header.paragraphs[0].text='PHSS | Evidence-linked research report'
    for kind,value in report_blocks(title,topic,sources,sections,**options):
        if kind=='heading':document.add_heading(clean(value[0]),value[1])
        elif kind=='paragraph':document.add_paragraph(clean(value))
        else:
            headers,rows=value
            table=document.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1'
            for cell,text in zip(table.rows[0].cells,headers):cell.text=clean(text)
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                for cell,text in zip(table.add_row().cells,row):
                    cell.text=clean(text)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:run.font.size=Pt(9)
            document.add_paragraph()
    footer=document.sections[0].footer.paragraphs[0]
    footer.alignment=2;footer.add_run('Page ')
    page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    out=BytesIO();document.save(out);return out.getvalue()

def readable_markdown(title,topic,sources,sections,**options):
    lines=[]
    for kind,value in report_blocks(title,topic,sources,sections,**options):
        if kind=='heading':lines.append('#'*max(1,value[1]+1)+' '+value[0])
        elif kind=='paragraph':lines.append(value)
        else:
            headers,rows=value
            def row(values):return '| '+' | '.join(str(v).replace('|','\\|').replace('\n','<br>') for v in values)+' |'
            lines+=[row(headers),row(['---']*len(headers))]+[row(values) for values in rows]
        lines.append('')
    return '\n'.join(lines)
