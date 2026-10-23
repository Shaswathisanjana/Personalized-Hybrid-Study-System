"""Report truthfulness tests: verified text is not verified interpretation."""
from io import BytesIO
from docx import Document
from app.research_agent.collection_models import FIELDS
from backend.research_reports import report_blocks,readable_docx,readable_markdown


def sources():
    def source(i,fields):
        claims={k:[] for k in FIELDS}
        for field,quote in fields.items():
            claims[field]=[dict(quote=quote,page=None,chunk_id='abstract-c1',document_id=str(i))]
        return dict(id=str(i),title=f'Synthetic source {i}',authors=[],year=None,doi=None,url='https://example.org/fixture',filename=None,scope='abstract_only',claims=claims,warnings=[],errors=[],total=1,completed=1,status='ready')
    return [source(1,{'models_algorithms':'Future directions include using Transformers.','dataset':'We discuss widely used datasets and assessment indices.','findings':'The method reports 98.3% mAP.','limitations':'Potential future directions are identified.'}),source(2,{'findings':'The review reports accuracy above 95%.'})]


def text(sections):return str(report_blocks('Fixture','Fixture topic',sources(),sections))


def test_future_work_is_not_presented_as_implemented_algorithm():
    t=text(['methods'])
    assert 'whether a method was implemented rather than reviewed or proposed' in t
    assert 'Future directions include using Transformers.' not in t
    assert 'does not verify scientific correctness' in t


def test_generic_dataset_and_different_metrics_are_not_claimed_comparable():
    t=text(['comparison','evaluation'])
    assert 'We discuss widely used datasets and assessment indices.' not in t
    assert 'Generic references to datasets do not identify a dataset' in t
    assert '98.3% mAP' in t and 'accuracy above 95%' in t
    assert 'no numerical ranking or cross-paper improvement is inferred' in t
    assert 'No best-performing method is established' in t


def test_generic_limitation_does_not_generate_invented_gap_or_experiment():
    t=text(['gaps','directions'])
    assert 'No suggestion is generated' in t
    assert 'not evidence of a research gap' in t
    assert 'verified novelty claims' in t
    assert 'Potential future directions are identified.' not in t


def test_summary_counts_saved_evidence_and_explains_omitted_sections():
    t=text(['overviews'])
    assert 'Findings: excerpts available for 2/2' in t
    assert 'Methodology: excerpts available for 0/2' in t
    assert 'Sections not requested' in t and 'Select all report sections' in t
    assert 'does not prove the paper lacks that information' in t


def test_valid_word_and_markdown_preserve_real_quotes_and_citations():
    options=['title','overviews','methods','evaluation','comparison','gaps','directions','limitations']
    blob=readable_docx('Fixture','Fixture topic',sources(),options)
    doc=Document(BytesIO(blob));content='\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.tables)==2 and '[1, abstract-only, chunk abstract-c1]' in content
    assert 'Future directions include using Transformers.' not in content
    assert 'Reported results and comparison boundaries' in content
    md=readable_markdown('Fixture','Fixture topic',sources(),options)
    assert '98.3% mAP' in md and 'not a common benchmark' in md
