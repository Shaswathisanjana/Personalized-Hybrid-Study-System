from io import BytesIO
from unittest.mock import patch
from docx import Document
from backend.tests.test_api import client, post, Runtime
from backend.service import Application
from app.research_agent.docx_exporter import SECTIONS
from pdf_collection_test_fixtures import make_pdf


def add(c,count=3):
    post(c,'/query',{'text':'Find research papers about agentic AI'})
    results=c.get(c.base+'/research').json()['search']['results']
    return post(c,'/research/add',{'ids':[s['id'] for s in results[:count]]})


def test_unanalyzed_sources_cannot_masquerade_as_analysis(client):
    r=add(client)
    assert r['report_readiness']['analyzed_chunks']==0
    response=client.post(client.base+'/research/report/docx',json={})
    assert response.status_code==409 and 'Prepare selected sources' in response.text
    assert not client.get(client.base).json()['evidence']


def test_explicit_bibliography_is_valid_and_honestly_labelled(client):
    add(client)
    response=client.post(client.base+'/research/report/docx',json={'metadata_only':True,'sections':['title','references','limitations']})
    assert response.status_code==200,response.text
    doc=Document(BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Source bibliography' in text and 'does not present completed research analysis' in text
    assert len(doc.tables)==1
    assert client.post(client.base+'/research/report/docx',json={'metadata_only':True,'sections':['methods']}).status_code==409


def test_preparation_then_readable_all_sections_docx_and_markdown(client):
    c=client;add(c)
    r=post(c,'/research/prepare-report');assert r['report_readiness']['ready']
    assert r['report_readiness']['analyzed_chunks']==3
    body={'title':'Controlled research report','topic':'Synthetic fixtures','sections':list(SECTIONS)}
    response=c.post(c.base+'/research/report/docx',json=body)
    assert response.status_code==200,response.text
    doc=Document(BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)
    for heading in SECTIONS.values():
        if heading not in (SECTIONS['title'],SECTIONS['references']):assert heading in text
    assert len(doc.tables)==2 and 'abstract-only' in text and 'chunk' in text
    assert 'The methodology uses a controlled comparison.' in text
    assert 'not reported in the analyzed source excerpts' in text
    assert 'System-generated research suggestions, not paper findings' in text
    assert 'https://example.org/offline-fixture/' in text
    md=c.post(c.base+'/research/report/md',json=body)
    assert md.status_code==200 and 'Methodology |' in md.text
    assert 'abstract-only' in md.text and 'chunk' in md.text
    assert not c.get(c.base).json()['evidence']
    fresh=Application(c.service.store.path,factory=Runtime)
    assert fresh.research_view(fresh.load(c.student))['report_readiness']['ready']
    assert not c.get('/api/students/'+c.other+'/research').json()['documents']


def test_seven_sources_bounded_partial_then_complete(client):
    c=client;add(c,7)
    first=post(c,'/research/prepare-report')['report_readiness']
    assert first['analyzed_chunks']==6 and not first['ready']
    assert c.post(c.base+'/research/report/docx',json={}).status_code==409
    partial=c.post(c.base+'/research/report/docx',json={'allow_partial':True})
    assert partial.status_code==200
    assert 'PARTIAL REPORT' in '\n'.join(p.text for p in Document(BytesIO(partial.content)).paragraphs)
    second=post(c,'/research/prepare-report')['report_readiness']
    assert second['ready'] and second['analyzed_chunks']==7
    assert post(c,'/research/prepare-report')['report_readiness']['analyzed_chunks']==7
    assert not c.get(c.base).json()['evidence']


def test_analysis_failure_recoverable_and_export_never_calls_provider(client):
    c=client;add(c,1)
    with patch('pdf_collection_test_fixtures.ExtractiveAnalyzer.analyze',side_effect=RuntimeError('offline failure')):
        r=post(c,'/research/prepare-report')['report_readiness']
    assert not r['ready'] and r['failed_chunks']==1
    assert c.post(c.base+'/research/report/docx',json={}).status_code==409
    assert post(c,'/research/prepare-report')['report_readiness']['ready']
    with patch('pdf_collection_test_fixtures.ExtractiveAnalyzer.analyze',side_effect=AssertionError('Export must not analyze')):
        assert c.post(c.base+'/research/report/docx',json={}).status_code==200


def test_selected_pdf_preparation_and_page_provenance(client):
    c=client
    for i in range(2):
        response=c.post(c.base+'/research/upload',files=[('files',(f'paper{i}.pdf',make_pdf(i),'application/pdf'))])
        assert response.status_code==200,response.text
    docs=response.json()['collection']['documents']
    response=c.post(c.base+'/research/select',json={'ids':[docs[1]['document_id']],'selected':False})
    assert response.status_code==200,response.text
    r=post(c,'/research/prepare-report')
    assert r['report_readiness']['selected_sources']==1 and r['report_readiness']['ready']
    assert next(d for d in r['documents'] if d['document_id']==docs[1]['document_id'])['record']['status']=='queued'
    report=c.post(c.base+'/research/report/md',json={})
    assert report.status_code==200 and 'page 2' in report.text and 'full-text excerpts' in report.text


def test_analysis_without_supported_claims_does_not_enable_report(client):
    c=client;add(c,1)
    with patch('pdf_collection_test_fixtures.ExtractiveAnalyzer.analyze',return_value={}):
        r=post(c,'/research/prepare-report')['report_readiness']
    assert not r['ready'] and r['supported_claims']==0
    assert c.post(c.base+'/research/report/docx',json={'allow_partial':True}).status_code==409


