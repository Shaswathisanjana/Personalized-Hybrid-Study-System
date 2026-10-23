"""Actual Word export of selected, grounded research content; no generation calls."""
from io import BytesIO
import re

from app.research_agent.collection_models import FIELDS
from app.research_agent.research_report import evidence_groups, location, research_suggestions

SECTIONS = {
    'title': 'Research topic and report title',
    'references': 'Analyzed paper details and references',
    'overviews': 'Individual paper overviews',
    'objectives': 'Research objectives',
    'methods': 'Methodologies',
    'evaluation': 'Dataset and evaluation information',
    'comparison': 'Methodology comparison table',
    'themes': 'Thematic literature synthesis',
    'differences': 'Similarities and differences',
    'gaps': 'Preliminary research gaps',
    'directions': 'Proposed future research directions',
    'limitations': 'Limitations and source coverage',
}


def clean(value):
    # XML 1.0 forbids control characters occasionally extracted from PDF text.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(value))


def build_docx(title, topic, sources, sections):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    sections = set(sections)
    if not sections or sections - set(SECTIONS):
        raise ValueError('Choose one or more supported report sections.')
    if not sources:
        raise ValueError('Select at least one source for the report.')
    document = Document()
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    document.sections[0].header.paragraphs[0].text = 'PHSS — evidence-linked research analysis'
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.75)
    def paragraph(text):
        document.add_paragraph(clean(text))
    def heading(text, level=1):
        document.add_heading(clean(text), level)
    def link(url):
        from urllib.parse import urlsplit
        if urlsplit(url).scheme not in ('http', 'https'):
            return
        para = document.add_paragraph()
        element = OxmlElement('w:hyperlink')
        element.set(qn('r:id'), para.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run, text = OxmlElement('w:r'), OxmlElement('w:t')
        text.text = clean(url)
        run.append(text)
        element.append(run)
        para._p.append(element)
    def evidence(source, field):
        for claim in source['claims'][field]:
            if claim.get('observation') and claim['observation'] != claim['quote']:
                paragraph('Saved source-supported analysis: ' + claim['observation'])
            paragraph(claim['quote'])
            paragraph('Source: ' + source['title'] + ' — ' + location(source, claim))
    if 'title' in sections:
        heading(title.strip() or 'Research Analysis Report', 0)
        paragraph('Research topic: ' + (topic.strip() or 'Not supplied'))
    # Coverage is a mandatory attribution notice, independent of optional sections.
    paragraph('Observations are source excerpts or labelled saved analyses accompanied by checked quotations. '
              'System suggestions are labelled separately. '
              'Abstract-only sources were not analyzed as full text. PDF coverage consists of extracted page chunks; '
              'it may be incomplete. This report is not a systematic literature review.')
    for source in sources:
        paragraph(f"Coverage — {source['title']}: {source['scope']}; "
                  f"{source['completed']}/{source['total']} chunks analyzed; status {source['status']}.")
    if 'references' in sections:
        heading(SECTIONS['references'])
        for source in sources:
            heading(source['title'], 2)
            paragraph('Authors: ' + (', '.join(source['authors']) or 'Not supplied'))
            paragraph('Year: ' + str(source['year'] or 'Not supplied'))
            paragraph('DOI: ' + (source['doi'] or 'Not supplied'))
            paragraph('Source URL: ' + (source['url'] or 'Local uploaded PDF; no publication URL supplied'))
            if source['url']:
                link(source['url'])
            paragraph('Source ID: ' + source['id'] + '; scope: ' + source['scope'])
            if source['filename']:
                paragraph('Original filename: ' + source['filename'])
        if sections - {'title', 'references'}:
            document.add_page_break()
    groups = dict(overviews=('overview', 'research_problem', 'findings'), objectives=('objectives',),
                  methods=('methodology', 'models_algorithms'),
                  evaluation=('dataset', 'experimental_setup', 'evaluation_metrics'))
    for key, fields in groups.items():
        if key not in sections:
            continue
        available = [s for s in sources if any(s['claims'][f] for f in fields)]
        if not available:
            continue  # Do not invent placeholder analyses for unavailable fields.
        heading(SECTIONS[key])
        for source in available:
            heading(source['title'] + ' — ' + source['scope'], 2)
            for field in fields:
                if source['claims'][field]:
                    heading(field.replace('_', ' ').title(), 3)
                    evidence(source, field)
    if 'comparison' in sections and any(s['claims']['methodology'] for s in sources):
        heading(SECTIONS['comparison'])
        table = document.add_table(rows=1, cols=3)
        table.style = 'Light Shading Accent 1'
        for cell, name in zip(table.rows[0].cells, ('Paper / source scope', 'Methodology excerpt', 'Attribution')):
            cell.text = name
        for source in sources:
            for claim in source['claims']['methodology']:
                cells = table.add_row().cells
                for cell, value in zip(cells, (source['title'] + '\n' + source['scope'], claim['quote'], location(source, claim))):
                    cell.text = clean(value)
    grouped = evidence_groups(sources)
    if 'themes' in sections and any(grouped.values()):
        heading(SECTIONS['themes'])
        paragraph('Evidence is synthesized by reported research field. Group membership does not establish '
                  'semantic equivalence, consensus or scientific validity.')
        for field, entries in grouped.items():
            if entries:
                heading(field.replace('_', ' ').title(), 2)
                for source, claim in entries:
                    paragraph(source['title'] + ': ' + claim['quote'])
                    paragraph(location(source, claim))
    if 'differences' in sections and any(grouped.values()):
        heading(SECTIONS['differences'])
        paragraph('Matching or distinct quotation wording is an observation about the sources, '
                  'not proof of identical methods or contradictory results.')
        for field, entries in grouped.items():
            wording = {}
            for source, claim in entries:
                wording.setdefault(' '.join(claim['quote'].casefold().split()), []).append((source, claim))
            if not wording:
                continue
            heading(field.replace('_', ' ').title(), 2)
            paragraph(f'{len(wording)} distinct reported descriptions across {len({s["id"] for s,c in entries})} sources.')
            for items in wording.values():
                paragraph(('Matching wording: ' if len({s['id'] for s,c in items}) > 1 else 'Reported wording: ') + items[0][1]['quote'])
                for source, claim in items:
                    paragraph(source['title'] + ' — ' + location(source, claim))
    limitations = grouped['limitations']
    gap_report, proposal = research_suggestions(topic, sources)
    for key in ('gaps', 'directions'):
        if key not in sections or not limitations:
            continue
        heading(SECTIONS[key])
        paragraph('AI-system suggestions derived from reported limitations; these are not paper findings '
                  'or verified gaps. Broader literature review and feasibility checks are required.')
        if key == 'gaps':
            for candidate in gap_report.candidates:
                paragraph(candidate.opportunity)
                paragraph(candidate.paper_title + ' — ' + candidate.paper_url)
        else:
            for direction in proposal.directions:
                paragraph(direction.objective)
                paragraph(direction.source_paper + ' — ' + direction.source_url)
    if 'limitations' in sections:
        heading(SECTIONS['limitations'])
        for source in sources:
            paragraph(f"{source['title']} — {source['scope']}: {source['completed']}/{source['total']} chunks analyzed; status {source['status']}.")
            for warning in source['warnings'] + source['errors']:
                paragraph(warning)
            evidence(source, 'limitations')
    output = BytesIO()
    document.save(output)
    return output.getvalue()
