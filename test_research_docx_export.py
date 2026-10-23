"""Read generated Word packages back with python-docx; all evidence is synthetic."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from app.persistence.sqlite_store import SQLiteStudentStore
from app.persistence.research_library_store import UnifiedResearchDocumentStore
from app.research_agent.collection_service import ResearchCollectionService
from app.research_agent.docx_exporter import build_docx, SECTIONS
from app.research_agent.research_report import report_sources
from pdf_collection_test_fixtures import make_pdf, ExtractiveAnalyzer
from test_research_search_collection import synthetic_paper


class DocxResearchReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = SQLiteStudentStore(Path(self.temp.name) / 'students.sqlite3')
        self.owner = self.store.create_profile('Synthetic report owner')
        pdfs = ResearchCollectionService(self.store, self.owner)
        pdfs.add('original-fixture.pdf', make_pdf())
        pdfs.extract_batch()
        self.library = UnifiedResearchDocumentStore(self.store)
        self.library.add_paper(self.owner, synthetic_paper(1))
        collection = ResearchCollectionService(self.store, self.owner, document_store=self.library,
                                               analyzer=ExtractiveAnalyzer())
        collection.analyze_batch()
        self.sources = report_sources(collection.list())

    @staticmethod
    def text(document):
        return '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells)

    def test_valid_complete_docx_has_actual_quotes_scopes_refs_and_table(self):
        output = build_docx('Selected Research Report', 'Synthetic topic', self.sources, SECTIONS)
        with ZipFile(BytesIO(output)) as package:
            self.assertIn('word/document.xml', package.namelist())
            self.assertIsNone(package.testzip())
        document = Document(BytesIO(output))
        text = self.text(document)
        self.assertIn('Selected Research Report', text)
        self.assertIn('The methodology uses a controlled comparison.', text)
        self.assertIn('full_text_excerpts', text)
        self.assertIn('abstract_only', text)
        self.assertIn('page 2', text)
        self.assertIn('abstract-c1', text)
        self.assertIn('original-fixture.pdf', text)
        self.assertTrue(document.tables)
        self.assertIn(self.sources[1]['url'], text)
        self.assertTrue(any(rel.is_external and rel.target_ref == self.sources[1]['url']
                            for rel in document.part.rels.values()))

    def test_sections_and_sources_are_selected_without_other_analysis_leaking(self):
        output = build_docx('Methods only', 'Synthetic topic', [self.sources[0]], ['title', 'methods'])
        text = self.text(Document(BytesIO(output)))
        self.assertIn('Methodologies', text)
        self.assertNotIn('Dataset and evaluation information', text)
        self.assertNotIn('The dataset contains synthetic records.', text)
        self.assertNotIn(self.sources[1]['title'], text)
        self.assertNotIn('Preliminary research gaps', text)
        self.assertIn('page 2', text)

    def test_missing_analysis_is_omitted_and_never_replaced_with_findings(self):
        source = self.sources[0]
        for field in source['claims']:
            source['claims'][field] = []
        text = self.text(Document(BytesIO(build_docx('No claims', '', [source], SECTIONS))))
        self.assertNotIn('The findings report an improvement', text)
        self.assertNotIn('Methodologies', text)
        self.assertNotIn('Preliminary research gaps', text)
        self.assertIn('Limitations and source coverage', text)

    def test_suggestions_are_labelled_separately_and_keep_limitation_attribution(self):
        text = self.text(Document(BytesIO(build_docx('Report', '', self.sources, ['gaps', 'directions']))))
        self.assertIn('AI-system suggestions', text)
        self.assertIn('not paper findings', text)
        self.assertIn('The study is limited to synthetic examples.', text)
        self.assertIn('page 2', text)
        self.assertIn('abstract-c1', text)

    def test_partial_coverage_and_provider_failures_are_visible(self):
        self.sources[0]['completed'] = 1
        self.sources[0]['warnings'].append('Fixture image page has no OCR text.')
        self.sources[0]['errors'].append('Injected provider failure')
        text = self.text(Document(BytesIO(build_docx('Partial', '', self.sources, ['limitations']))))
        self.assertIn('1/2 chunks analyzed', text)
        self.assertIn('no OCR text', text)
        self.assertIn('Injected provider failure', text)

    def test_invalid_selection_is_rejected_and_xml_controls_are_sanitized(self):
        with self.assertRaises(ValueError):
            build_docx('Report', '', self.sources, [])
        with self.assertRaises(ValueError):
            build_docx('Report', '', [], ['title'])
        with self.assertRaises(ValueError):
            build_docx('Report', '', self.sources, ['invented-section'])
        output = build_docx('Report\x00\x0b', '', self.sources, ['title'])
        self.assertIn('Report', self.text(Document(BytesIO(output))))

    def test_existing_markdown_export_remains_available(self):
        from app.research_agent.report_exporter import ResearchReportExporter
        from app.research_agent.research_workflow import ResearchWorkflowResult
        text = ResearchReportExporter.build_markdown(ResearchWorkflowResult(topic='Original abstract workflow'))
        self.assertIn('Original abstract workflow', text)
        self.assertIn('abstract', text.lower())
        from app.research_agent.research_report import workflow_report_sources
        from app.research_agent.paper_summarizer import PaperSummary
        paper = synthetic_paper(3)
        quote = 'The methodology uses a controlled comparison.'
        result = ResearchWorkflowResult(topic='Saved analysis', accepted_papers=[paper],
                    summaries=[PaperSummary(paper.title, paper.authors, paper.year, paper.url,
                               proposed_method='A controlled comparison is described.',
                               evidence_quotes={'proposed_method': quote})])
        saved = workflow_report_sources(result)
        exported = self.text(Document(BytesIO(build_docx('Existing work', result.topic, saved, ['methods']))))
        self.assertIn('A controlled comparison is described.', exported)
        self.assertIn(quote, exported)
        self.assertIn('abstract_only', exported)


if __name__ == '__main__':
    unittest.main()
