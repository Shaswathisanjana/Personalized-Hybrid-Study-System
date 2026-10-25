"""Exercise original Unified document layout with external services isolated."""
import io
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from app.agents.research.word_export import export_word_report

class UnifiedExportTests(unittest.TestCase):
    def test_original_report_layout(self):
        config=types.ModuleType("src.config")
        config.GROQ_MODEL="test"
        config.OUTPUT_DIR=Path(tempfile.gettempdir())
        analysis=types.ModuleType("src.paper_analysis")
        analysis.MISSING="Not explicitly reported in full text"
        analysis._clean_json_response=lambda text:text
        analysis._is_noisy_summary=lambda *args:False
        analysis._looks_unsupported_or_fragmentary=lambda *args:False
        dotenv=types.ModuleType("dotenv");dotenv.load_dotenv=lambda *a,**k:None
        report={"topic":"Unified fixture","paper_analyses":[{
            "title":"Evidence fixture","year":"2025","authors":["Test Author"],
            "methodology":"Train and evaluate on held-out samples.",
            "dataset":"Fixture data","results_metrics":"90% fixture accuracy",
            "author_stated_limitations":"Small sample size.",
            "potential_gap_from_evidence":"Test larger samples.",
            "url":"https://example.com/paper"}]}
        with patch.dict(sys.modules,{"src.config":config,"src.paper_analysis":analysis,"dotenv":dotenv}),patch.dict("os.environ",{"GROQ_API_KEY":""}):
            payload=export_word_report(report)
        doc=Document(io.BytesIO(payload))
        text="\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Research Literature Review Report",text)
        self.assertIn("2. Paper-wise comparison",text)
        self.assertGreater(len(doc.tables),0)
        cells=" ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn("Evidence fixture",cells)
        self.assertIn("90% fixture accuracy",cells)
        self.assertNotIn("PHSS · Research report",text)

if __name__=="__main__": unittest.main()
