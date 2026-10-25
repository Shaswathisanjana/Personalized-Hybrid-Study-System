import io,unittest,zipfile
from docx import Document
from app.agents.research.word_export import export_word_report

class WordExportTests(unittest.TestCase):
    def test_generates_real_word_file_without_ai_or_cached_paths(self):
        data={'topic':'Glaucoma detection','word_report_path':'/must/not/be/read.docx','papers':[{'title':'Sample paper','url':'https://example.org/paper','summary':'Provided abstract only.'}],'paper_analyses':[{'title':'Sample paper','methodology':'U-Net','datasets':['Provided dataset'],'results_metrics':'Reported value','limitations':'Small sample'}],'gaps':[{'title':'External validation','rationale':'Limited validation data'}]}
        out=export_word_report(data)
        with zipfile.ZipFile(io.BytesIO(out)) as z:self.assertIn('word/document.xml',z.namelist())
        doc=Document(io.BytesIO(out));alltext=' '.join(p.text for p in doc.paragraphs)+' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('Glaucoma detection',alltext);self.assertIn('U-Net',alltext);self.assertIn('External validation',alltext)
    def test_export_preserves_unicode_and_labels_missing_fields(self):
        doc=Document(io.BytesIO(export_word_report({'topic':'Study α','papers':[{'title':'Paper β'}]})))
        self.assertIn('Study α',' '.join(p.text for p in doc.paragraphs));self.assertTrue(any('Not reported' in c.text for t in doc.tables for r in t.rows for c in r.cells))
    def test_no_papers_is_still_a_valid_document(self):
        self.assertTrue(export_word_report({'topic':'Empty report'}).startswith(b'PK'))
if __name__=='__main__':unittest.main()
