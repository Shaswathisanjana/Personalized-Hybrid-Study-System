"""Export the same research-progression narrative displayed in the web app."""
import io
from docx import Document
from docx.shared import Pt, Inches

def export_progression(report):
    document=Document()
    document.styles["Normal"].font.name="Calibri"
    document.styles["Normal"].font.size=Pt(11)
    section=document.sections[0]
    section.top_margin=section.bottom_margin=Inches(.75)
    section.left_margin=section.right_margin=Inches(.8)
    normal=document.styles["Normal"].paragraph_format
    normal.space_after=Pt(8)
    normal.line_spacing=1.15
    document.add_heading("Research Literature Review and Progression",0)
    document.add_paragraph(report.get("topic",""))
    coverage=report.get("coverage") or {}
    document.add_heading("Search coverage and selection",1)
    document.add_paragraph("Searched at: "+str(coverage.get("searched_at","Unknown")))
    document.add_paragraph("Latest included publication: "+str(coverage.get("latest_included_publication","Unknown")))
    document.add_paragraph(str(report.get("selection_reason","")))
    document.add_paragraph(str(coverage.get("limitations","Bounded review; not exhaustive.")))
    summary=report.get("research_progression") or {}
    if summary.get("section_warnings"):
        document.add_heading("Incomplete sections",1)
        for warning in summary["section_warnings"]:document.add_paragraph(str(warning))
    for key,label in [("timeline","How the research progressed"),("current_state","Where the field stands"),
                      ("comparisons","Cross-paper technical comparison"),("gaps","Gaps in the reviewed evidence"),("future_work","Proposed future work")]:
        document.add_heading(label,1)
        for item in summary.get(key,[]):
            document.add_heading((item.get("period","")+" "+item.get("title","")).strip(),2)
            for paragraph in item.get("description","").split("\n"):
                if paragraph.strip():document.add_paragraph(paragraph.strip())
            for key,value in (item.get("experiment") or {}).items():
                document.add_heading(key.replace("_"," ").capitalize(),3)
                document.add_paragraph(str(value))
            document.add_paragraph("Supporting sources: "+", ".join(item.get("source_ids",[])))
    if report.get("paper_analyses"):
        document.add_heading("Supporting paper-wise evidence",1)
        document.add_paragraph("These records support the cross-paper synthesis. Missing fields are not inferred. Surveys summarize other studies and do not necessarily introduce their own dataset or experiment.")
        for paper in report["paper_analyses"]:
            document.add_heading(paper.get("title","Untitled"),2)
            for key,label in [("methodology","Approach"),("dataset","Data / evaluation setting"),("results_metrics","Reported findings"),("author_stated_limitations","Limitations")]:
                document.add_paragraph(label,style="Heading 3")
                document.add_paragraph(str(paper.get(key) or "Not established from available evidence."))
    document.add_heading("References",1)
    for p in report.get("papers",[]):
        document.add_paragraph("["+p.get("source_id","")+"] "+p.get("title","")+" ("+str(p.get("published") or p.get("year") or "Unknown date")[:10]+"). "+p.get("url",""))
    stream=io.BytesIO();document.save(stream);return stream.getvalue()
